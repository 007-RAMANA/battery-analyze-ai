
from pathlib import Path
import json
import re

import pandas as pd


# ---------------------------------------------------------
# BATTERY AI - FILE READER AND ANALYSIS ENGINE
# Supports CSV, TXT, LOG, MD, XLSX, XLS, PDF, DOCX, JSON
# Includes Japanese battery-log header detection
# ---------------------------------------------------------


def normalize_name(value):
    """Normalize column names while preserving Japanese and other scripts."""
    text = str(value).strip().casefold()
    text = text.replace("℃", "c").replace("°c", "c")
    text = text.replace("（", "(").replace("）", ")")
    text = text.replace("％", "%")

    # Remove spaces and common separators, but preserve Unicode letters.
    text = re.sub(r"[\s_\-./\\()[\]{}:%]+", "", text)
    return text


COLUMN_ALIASES = {

    "Voltage": [
        "voltage", "voltagev", "volt", "volts", "v",
        "電圧", "電圧v", "端子電圧", "電池電圧", "전압", "전압v",
    ],
    "Current": [
        "current", "currenta", "amps", "ampere", "amperes", "a",
        "電流", "電流a", "전류", "전류a",
    ],
    "Temperature": [
        "temperature", "temperaturec", "temp", "celltemperature",
        "celltemp", "temperature℃", "セル温度", "温度", "温度c",
        "周囲温度","セル温度c","셀온도c", "셀온도", "온도", "온도c",
    ],
    "Cycle": [
    "cycle", "cyclenumber", "cycleno", "cyclen",
    "サイクル", "サイクル数", "サイクル番号", "サイクルno",
    "사이클수", "사이클번호", "사이클",
],
    "Phase": [
    "phase", "step", "mode", "status", "operation", "teststep",
    "ステップ", "試験ステップ", "工程", "状態", "モード",
    "充放電", "試験モード",
    "시험단계",
],
    "Time": [
        "time", "timestamp", "datetime", "date", "elapsedtime",
        "elapsedtimes", "elapsedtime_seconds", "経過時間",
        "経過時間秒", "測定日時", "日時", "時刻", "측정일시", "경과시간초", "경과시간",
    ],
    "Charge Capacity": [
        "chargecapacity", "chargecapacityah", "chargingcapacity",
        "chargeah", "充電容量", "充電容量ah", "充電容量a h", "충전용량", "충전용량ah",
    ],
    "Discharge Capacity": [
        "dischargecapacity", "dischargecapacityah",
        "dischargingcapacity", "dischargeah",
        "放電容量", "放電容量ah","방전용량", "방전용량ah",
    ],
}


def _alias_lookup():
    lookup = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        lookup[normalize_name(canonical)] = canonical
        for alias in aliases:
            lookup[normalize_name(alias)] = canonical
    return lookup


ALIAS_LOOKUP = _alias_lookup()


def _canonicalize_columns(df):
    """Rename recognized column headers to standard English names."""
    if df is None or df.empty:
        return df

    new_columns = []
    used = set()

    for index, column in enumerate(df.columns):
        original = str(column).strip()
        normalized = normalize_name(original)
        canonical = ALIAS_LOOKUP.get(normalized, original)

        # Avoid duplicate column names after normalization.
        if canonical in used:
            canonical = f"{canonical}_{index + 1}"

        new_columns.append(canonical)
        used.add(canonical)

    result = df.copy()
    result.columns = new_columns
    return result


def _measurement_header_score(row):
    """Score a raw Excel row to find the actual measurement header."""
    values = [
        normalize_name(value)
        for value in row.tolist()
        if pd.notna(value) and str(value).strip()
    ]

    recognized = set()

    for value in values:
        canonical = ALIAS_LOOKUP.get(value)
        if canonical:
            recognized.add(canonical)

    # A battery measurement header should contain several known fields.
    return len(recognized)


def read_excel_battery_file(file_path, sheet_name=0):
    """
    Read Excel files with or without metadata rows before the table.
    For the Japanese sample, this detects the measurement header on row 9.
    """
    file_path = str(file_path)

    try:
        raw = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            header=None,
            dtype=object,
        )
    except Exception as exc:
        raise ValueError(f"Unable to read Excel file: {exc}") from exc

    if raw.empty:
        return pd.DataFrame()

    header_row = None

    # Find a row containing multiple recognized measurement headings.
    for row_index in range(len(raw)):
        if _measurement_header_score(raw.iloc[row_index]) >= 3:
            header_row = row_index
            break

    if header_row is not None:
        headers = []
        for index, value in enumerate(raw.iloc[header_row].tolist()):
            if pd.isna(value) or not str(value).strip():
                headers.append(f"Unnamed_{index}")
            else:
                headers.append(str(value).strip())

        data = raw.iloc[header_row + 1:].copy()
        data.columns = headers
        data = data.dropna(how="all").reset_index(drop=True)
        return _canonicalize_columns(data)

    # Fallback for ordinary Excel files whose first row is the header.
    try:
        ordinary = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            dtype=object,
        )
        return _canonicalize_columns(ordinary)
    except Exception as exc:
        raise ValueError(f"Unable to identify Excel table headers: {exc}") from exc


def _read_text_file(file_path):
    """Read common text formats using a few likely encodings."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return Path(file_path).read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    return Path(file_path).read_text(encoding="utf-8", errors="replace")


def _text_to_dataframe(text):
    """Try delimiter-separated data first, then fall back to plain text."""
    from io import StringIO

    text = text.strip()
    if not text:
        return pd.DataFrame()

    # Try common separators.
    for separator in (",", "\t", ";", "|"):
        try:
            candidate = pd.read_csv(StringIO(text), sep=separator)
            if len(candidate.columns) > 1:
                return _canonicalize_columns(candidate)
        except Exception:
            pass

    # Try whitespace-delimited tables.
    try:
        candidate = pd.read_csv(StringIO(text), sep=r"\s+", engine="python")
        if len(candidate.columns) > 1:
            return _canonicalize_columns(candidate)
    except Exception:
        pass

    # Last attempt: one-column text file.
    return pd.DataFrame({"Text": text.splitlines()})


def _read_pdf(file_path):
    """Extract tables from PDF when possible, otherwise extract text."""
    try:
        import fitz
    except ImportError as exc:
        raise ImportError(
            "PDF support requires PyMuPDF. Install it with: "
            "python -m pip install pymupdf"
        ) from exc

    collected_tables = []
    collected_text = []

    with fitz.open(str(file_path)) as document:
        for page in document:
            page_text = page.get_text("text")
            if page_text.strip():
                collected_text.append(page_text)

            try:
                tables = page.find_tables()
                for table in tables.tables:
                    extracted = table.extract()
                    if extracted and len(extracted) >= 2:
                        table_df = pd.DataFrame(
                            extracted[1:],
                            columns=extracted[0],
                        )
                        collected_tables.append(table_df)
            except Exception:
                # Some PDFs contain no extractable tables.
                pass

    if collected_tables:
        combined = pd.concat(collected_tables, ignore_index=True)
        return _canonicalize_columns(combined)

    return _text_to_dataframe("\n".join(collected_text))


def _read_docx(file_path):
    """Read paragraphs and tables from DOCX."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "DOCX support requires python-docx. Install it with: "
            "python -m pip install python-docx"
        ) from exc

    document = Document(str(file_path))
    tables = []

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) >= 2:
            tables.append(
                pd.DataFrame(rows[1:], columns=rows[0])
            )

    if tables:
        return _canonicalize_columns(
            pd.concat(tables, ignore_index=True)
        )

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    return _text_to_dataframe("\n".join(paragraphs))


def read_battery_file(file_path, sheet_name=0):
    """
    Read a battery data file into a DataFrame.
    sheet_name may be an Excel sheet name or sheet index.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    extension = path.suffix.lower()

    if extension in (".xlsx", ".xls"):
        return read_excel_battery_file(path, sheet_name=sheet_name)

    if extension == ".csv":
        errors = []
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return _canonicalize_columns(
                    pd.read_csv(path, encoding=encoding)
                )
            except UnicodeDecodeError as exc:
                errors.append(str(exc))
        raise ValueError("Unable to decode CSV file.")

    if extension in (".txt", ".log", ".md"):
        return _text_to_dataframe(_read_text_file(path))

    if extension == ".json":
        try:
            with path.open("r", encoding="utf-8-sig") as handle:
                content = json.load(handle)

            if isinstance(content, list):
                df = pd.DataFrame(content)
            elif isinstance(content, dict):
                # Handle common JSON wrappers around record arrays.
                records = None
                for key in ("data", "records", "measurements", "results"):
                    if isinstance(content.get(key), list):
                        records = content[key]
                        break

                if records is not None:
                    df = pd.DataFrame(records)
                else:
                    df = pd.json_normalize(content)
            else:
                df = pd.DataFrame({"Value": [content]})

            return _canonicalize_columns(df)
        except Exception as exc:
            raise ValueError(f"Unable to parse JSON file: {exc}") from exc

    if extension == ".pdf":
        return _read_pdf(path)

    if extension == ".docx":
        return _read_docx(path)

    raise ValueError(
        f"Unsupported file type: {extension}. "
        "Supported: CSV, TXT, LOG, MD, XLSX, XLS, PDF, DOCX, JSON."
    )


def _numeric_series(df, column):
    """Convert a column to numeric values, ignoring unparseable cells."""
    if column not in df.columns:
        return pd.Series(dtype="float64")

    series = df[column]

    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce").dropna()

    cleaned = (
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("％", "%", regex=False)
        .str.replace("%", "", regex=False)
        .str.strip()
    )
    return pd.to_numeric(cleaned, errors="coerce").dropna()


def _last_numeric_value(df, column):
    values = _numeric_series(df, column)
    if values.empty:
        return None
    return float(values.iloc[-1])


def _first_present_column(df, names):
    for name in names:
        if name in df.columns:
            return name
    return None


def analyze_battery_data(df):
    """
    Analyze a DataFrame and return dashboard-compatible results:
    rows, columns, metrics, warnings, phase_summary,
    cycle_summary, and capacity_ratio.
    """
    if df is None:
        df = pd.DataFrame()

    df = df.copy()
    df = _canonicalize_columns(df)

    warnings = []
    metrics = {}

    rows = int(len(df))
    columns = [str(column) for column in df.columns]

    if df.empty:
        warnings.append("The selected file contains no data rows.")

    numeric_columns = {}
    for column in df.columns:
        values = _numeric_series(df, column)
        if not values.empty:
            numeric_columns[column] = values

    # Known battery metrics.
    voltage = _numeric_series(df, "Voltage")
    current = _numeric_series(df, "Current")
    temperature = _numeric_series(df, "Temperature")
    charge_capacity = _numeric_series(df, "Charge Capacity")
    discharge_capacity = _numeric_series(df, "Discharge Capacity")
    cycle = _numeric_series(df, "Cycle")

    metrics["latest_voltage"] = (
        float(voltage.iloc[-1]) if not voltage.empty else None
    )
    metrics["latest_current"] = (
        float(current.iloc[-1]) if not current.empty else None
    )
    metrics["latest_temperature"] = (
        float(temperature.iloc[-1]) if not temperature.empty else None
    )
    metrics["charge_capacity"] = (
        float(charge_capacity.max()) if not charge_capacity.empty else None
    )
    metrics["discharge_capacity"] = (
        float(discharge_capacity.max())
        if not discharge_capacity.empty else None
    )
    metrics["cycle_count"] = (
        int(cycle.max()) if not cycle.empty else None
    )

    # Alternative display keys for compatibility with dashboards.
    metrics["Voltage"] = metrics["latest_voltage"]
    metrics["Current"] = metrics["latest_current"]
    metrics["Temperature"] = metrics["latest_temperature"]
    metrics["Charge Capacity"] = metrics["charge_capacity"]
    metrics["Discharge Capacity"] = metrics["discharge_capacity"]
    metrics["Cycle"] = metrics["cycle_count"]

    capacity_ratio = None
    if (
        not charge_capacity.empty
        and not discharge_capacity.empty
        and charge_capacity.max() > 0
    ):
        capacity_ratio = float(
            discharge_capacity.max() / charge_capacity.max()
        )

    metrics["capacity_ratio"] = capacity_ratio

    if not any(
        not series.empty
        for series in (voltage, current, temperature,
                       charge_capacity, discharge_capacity)
    ):
        warnings.append(
            "No recognized numeric battery measurement columns were found. "
            "Check the file's table headers and units."
        )

    # Phase summary: count rows and summarize numeric readings per phase.
    phase_summary = {}
    if "Phase" in df.columns:
        phase_values = df["Phase"].fillna("Unknown").astype(str)
        for phase, indices in phase_values.groupby(phase_values).groups.items():
            subset = df.loc[indices]
            phase_summary[str(phase)] = {"rows": int(len(subset))}

            for column in (
                "Voltage", "Current", "Temperature",
                "Charge Capacity", "Discharge Capacity",
            ):
                values = _numeric_series(subset, column)
                if not values.empty:
                    phase_summary[str(phase)][column] = {
                        "min": float(values.min()),
                        "max": float(values.max()),
                        "mean": float(values.mean()),
                    }

    # Cycle summary as a list of dictionaries for the dashboard.
    cycle_summary = []
    if "Cycle" in df.columns:
        work = df.copy()
        work["_cycle_numeric"] = pd.to_numeric(
            work["Cycle"], errors="coerce"
        )
        work = work.dropna(subset=["_cycle_numeric"])

        for cycle_number, subset in work.groupby("_cycle_numeric"):
            item = {"Cycle": int(cycle_number), "Rows": int(len(subset))}

            for column in (
                "Voltage", "Current", "Temperature",
                "Charge Capacity", "Discharge Capacity",
            ):
                values = _numeric_series(subset, column)
                if not values.empty:
                    item[f"Max {column}"] = float(values.max())
                    item[f"Min {column}"] = float(values.min())

            cycle_summary.append(item)

    return {
        "rows": rows,
        "columns": columns,
        "metrics": metrics,
        "warnings": warnings,
        "phase_summary": phase_summary,
        "cycle_summary": cycle_summary,
        "capacity_ratio": capacity_ratio,
        "dataframe": df,
        "numeric_columns": list(numeric_columns.keys()),
    }
